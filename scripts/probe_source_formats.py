"""Upload real synthetic format fixtures through the local parser API."""

from __future__ import annotations

import argparse
import asyncio
import io
import json
import zipfile
from pathlib import Path
from urllib.parse import urlparse

import httpx
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas


def fixtures():
    text = "SYNTHETIC TEST FILE. Need twelve valves with HSE certificates. No order placed."
    files = {
        "brief.txt": text.encode(),
        "brief.md": text.encode(),
        "brief.json": json.dumps({"brief": text}).encode(),
        "quotes.csv": (
            b"quote_id,item,quantity,unit_price,delivery,certified\n"
            b"A,Valve,12,120000,80000,true\nC,Valve,12,130000,0,true\n"
        ),
    }
    doc = Document()
    doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    files["brief.docx"] = buffer.getvalue()
    book = Workbook()
    book.active.append(["Brief", text])
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    files["brief.xlsx"] = buffer.getvalue()
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.drawString(40, 750, text)
    pdf.save()
    files["brief.pdf"] = buffer.getvalue()
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("brief.docx", files["brief.docx"])
        archive.writestr("brief.txt", files["brief.txt"])
    files["packet.zip"] = buffer.getvalue()
    image = Image.new("RGB", (1600, 260), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(
        str(
            Path(__file__).resolve().parents[1] / "src/ai_orchestrator/report_assets/DejaVuSans.ttf"
        ),
        32,
    )
    draw.text((30, 55), "SYNTHETIC OCR TEST. Need twelve valves.", fill="black", font=font)
    draw.text((30, 120), "HSE certificates required. No order placed.", fill="black", font=font)
    for extension, format in [("png", "PNG"), ("jpeg", "JPEG"), ("tiff", "TIFF")]:
        buffer = io.BytesIO()
        image.save(buffer, format=format)
        files["scan." + extension] = buffer.getvalue()
    buffer = io.BytesIO()
    image.save(buffer, format="PDF")
    files["scanned.pdf"] = buffer.getvalue()
    return files


async def main(args):
    if urlparse(args.base).hostname not in {"localhost", "127.0.0.1"}:
        raise ValueError("Local API only")
    args.output.mkdir(parents=True, exist_ok=True)
    report = {"synthetic": True, "cases": [], "passed": False}
    async with httpx.AsyncClient(
        base_url=args.base, headers={"x-organization-id": args.org}, timeout=60
    ) as client:
        for filename, content in fixtures().items():
            (args.output / filename).write_bytes(content)
            response = await client.post(
                "/api/v1/source-files", files={"file": (filename, content)}
            )
            response.raise_for_status()
            file = response.json()
            detail = (await client.get("/api/v1/source-files/" + file["id"])).json()
            extraction = detail["extraction"]
            case = {
                "filename": filename,
                "state": file["state"],
                "source_id": file["id"],
                "segments": len(extraction["segments"]),
                "warnings": extraction["warnings"],
                "error": extraction.get("error"),
            }
            case["passed"] = file["state"] in {"ready", "needs_review"} and case["segments"] > 0
            report["cases"].append(case)
    report["passed"] = all(c["passed"] for c in report["cases"])
    (args.output / "acceptance.json").write_text(json.dumps(report, indent=2))
    print(json.dumps(report))
    if not report["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base", default="http://127.0.0.1:8100")
    parser.add_argument("--org", required=True)
    parser.add_argument("--output", type=Path, required=True)
    asyncio.run(main(parser.parse_args()))
