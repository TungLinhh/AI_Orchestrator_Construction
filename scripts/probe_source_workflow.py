"""Exercise local source intake and real-model execution; verify downloaded file contents."""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import io
import json
import time
import zipfile
from pathlib import Path
from urllib.parse import urlparse

import httpx
from docx import Document
from openpyxl import load_workbook


async def main(args):
    if urlparse(args.base).hostname not in {"localhost", "127.0.0.1"}:
        raise ValueError("Local API only")
    args.output.mkdir(parents=True, exist_ok=True)
    report = {"synthetic": True, "passed": False, "cases": []}
    receipt = args.output / "acceptance.json"

    def save():
        receipt.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")

    async with httpx.AsyncClient(
        base_url=args.base, headers={"x-organization-id": args.org}, timeout=60
    ) as client:

        async def call(method, path, **kwargs):
            response = await client.request(method, "/api/v1" + path, **kwargs)
            response.raise_for_status()
            return response.json()

        agents = await call("GET", "/agents?limit=200")
        owner = next(a for a in agents["items"] if a["name"] == "Procurement Agent")
        for mode in ["mock", "sources"]:
            case = {"mode": mode, "passed": False}
            report["cases"].append(case)
            save()
            try:
                workflow = "quote_comparison" if mode == "mock" else "document_review"
                if mode == "mock":
                    files = (await call("POST", "/source-files/mock?workflow=quote_comparison"))[
                        "items"
                    ]
                else:
                    doc = Document()
                    doc.add_paragraph(
                        "SYNTHETIC ACCEPTANCE: Hồ sơ van cần chứng chỉ HSE. Chưa đặt hàng."
                    )
                    buffer = io.BytesIO()
                    doc.save(buffer)
                    packet = io.BytesIO()
                    with zipfile.ZipFile(packet, "w") as archive:
                        archive.writestr("brief.docx", buffer.getvalue())
                    file = await call(
                        "POST",
                        "/source-files",
                        files={"file": ("acceptance.zip", packet.getvalue())},
                    )
                    assert file["state"] == "ready", file
                    files = [file]
                case["source_ids"] = [f["id"] for f in files]
                payload = {
                    "title": "[SYNTHETIC SOURCE ACCEPTANCE] " + mode,
                    "goal": (
                        "Review selected documents with exact citations. "
                        "Identify missing certificates. No external actions."
                    ),
                    "workflow": workflow,
                    "owner_agent_id": owner["id"],
                    "selection": {"mode": mode, "file_ids": case["source_ids"]},
                }
                task = await call("POST", "/source-workflows", json=payload)
                case["task_id"] = task["id"]
                save()
                case["run"] = await call("POST", f"/tasks/{task['id']}/run")
                deadline = time.monotonic() + args.deadline
                while time.monotonic() < deadline:
                    task = await call("GET", f"/tasks/{task['id']}")
                    if task["status"] in {
                        "completed",
                        "failed",
                        "canceled",
                        "waiting_for_input",
                        "waiting_for_approval",
                    }:
                        break
                    await asyncio.sleep(3)
                case["status"] = task["status"]
                case["error"] = task.get("last_error")
                assert task["status"] == "completed", case
                artifacts = (await call("GET", f"/tasks/{task['id']}/artifacts"))["items"]
                assert {a["format"] for a in artifacts} == {
                    "docx",
                    "xlsx",
                    "pdf",
                    "json",
                    "md",
                    "zip",
                }
                case["artifacts"] = artifacts
                contents = {}
                for artifact in artifacts:
                    response = await client.get(
                        "/api/v1/source-artifacts/" + artifact["id"] + "/download"
                    )
                    response.raise_for_status()
                    body = response.content
                    assert hashlib.sha256(body).hexdigest() == artifact["sha256"]
                    (args.output / (mode + "." + artifact["format"])).write_bytes(body)
                    contents[artifact["format"]] = body
                assert len(Document(io.BytesIO(contents["docx"])).paragraphs) > 3
                book = load_workbook(io.BytesIO(contents["xlsx"]))
                assert book.active.max_row > 3
                book.close()
                assert contents["pdf"].startswith(b"%PDF-")
                result = json.loads(contents["json"])
                case["output"] = result
                if mode == "mock":
                    assert (
                        result["output"]["winner"] == "A"
                        and result["output"]["total_vnd"] == 1520000
                    )
                with zipfile.ZipFile(io.BytesIO(contents["zip"])) as archive:
                    assert archive.testzip() is None
                    assert any(name.startswith("sources/") for name in archive.namelist())
                case["passed"] = True
            except Exception as exc:
                case["exception"] = str(exc)[:1500]
            save()
    report["passed"] = all(c["passed"] for c in report["cases"])
    save()
    print(
        json.dumps(
            {
                "passed": report["passed"],
                "cases": [
                    {k: c.get(k) for k in ("mode", "task_id", "status", "passed", "exception")}
                    for c in report["cases"]
                ],
            }
        )
    )
    if not report["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base", default="http://127.0.0.1:8100")
    parser.add_argument("--org", required=True)
    parser.add_argument("--deadline", type=int, default=240)
    parser.add_argument("--output", type=Path, required=True)
    asyncio.run(main(parser.parse_args()))
