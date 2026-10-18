"""Explicit live HR acceptance: free-form goal, uploaded source, cited downloadable report.

Only run against the local, counted real-provider API. This command creates a task
and spends its configured model budget. Source data must be authorized by the user.
Expected phrases are acceptance checks and are never included in model input.
"""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import time
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from playwright.async_api import async_playwright


async def main(args: argparse.Namespace) -> None:
    if urlparse(args.base).hostname not in {"127.0.0.1", "localhost"}:
        raise ValueError("Local API only")
    args.output.mkdir(parents=True, exist_ok=False)
    goal = args.goal.read_text()
    receipt = {
        "passed": False,
        "kind": "live_hr_ordinary_uploaded_source",
        "source_sha256": hashlib.sha256(args.source.read_bytes()).hexdigest(),
        "goal_sha256": hashlib.sha256(goal.encode()).hexdigest(),
        "page_errors": [],
        "states": [],
    }
    headers = {"x-organization-id": args.org}
    async with async_playwright() as driver:
        browser = await driver.chromium.launch(executable_path=args.chromium, args=["--no-sandbox"])
        page = await browser.new_page(viewport={"width": 1440, "height": 1100})
        page.on("pageerror", lambda error: receipt["page_errors"].append(str(error)))
        try:
            url = args.base + "/api/v1/ui?org=" + args.org
            await page.goto(url + "#/give", wait_until="domcontentloaded")
            await page.locator("#newTaskBtn").click()
            await page.wait_for_function(
                "[...document.querySelectorAll('#agent option')]"
                ".some(o=>o.textContent.includes('HR Agent'))"
            )
            owner = await page.locator("#agent option").evaluate_all(
                "options=>options.find(o=>o.textContent.includes('HR Agent')).value"
            )
            await page.locator("#scenario").select_option("")
            await page.locator("#sourceWorkflow").select_option("")
            await page.locator("#taskType").select_option("analysis")
            await page.locator("#agent").select_option(owner)
            await page.locator("#goal").fill(goal)
            await page.locator("#sourceChooseFiles").check()
            await page.locator("#sourceUploadFiles").set_input_files(str(args.source))
            async with page.expect_response(
                lambda response: (
                    response.url.endswith("/source-files") and response.request.method == "POST"
                )
            ) as pending:
                await page.locator("#sourceUpload").click()
            uploaded = await pending.value
            assert uploaded.status == 201, await uploaded.text()
            source = await uploaded.json()
            assert source["state"] == "ready", source["state"]
            receipt["source_id"] = source["id"]
            await page.wait_for_function(
                "document.getElementById('sourceSelectionCount').textContent.includes('1')"
            )
            await page.screenshot(path=str(args.output / "source-selection.png"), full_page=True)
            async with page.expect_response(
                lambda response: (
                    response.url.endswith("/source-workflows") and response.request.method == "POST"
                )
            ) as pending:
                await page.locator("#runBtn").click()
            created = await pending.value
            assert created.status == 201, await created.text()
            task = await created.json()
            receipt["task_id"] = task["id"]
            assert task["goal"] == goal.strip()
            assert task["owner_agent_id"] == owner
            assert task["input"]["source_workflow"] == "document_review"
            assert task["input"]["source_selection_mode"] == "sources"
            deadline = time.monotonic() + args.timeout
            while time.monotonic() < deadline:
                response = await page.request.get(
                    args.base + "/api/v1/tasks/" + task["id"], headers=headers
                )
                assert response.status == 200
                task = await response.json()
                if not receipt["states"] or receipt["states"][-1] != task["status"]:
                    receipt["states"].append(task["status"])
                if task["status"] in {"completed", "failed", "cancelled", "waiting_for_input"}:
                    break
                await asyncio.sleep(2)
            assert task["status"] == "completed", task.get("last_error", task["status"])
            output = task["output"]
            report_text = json.dumps(output, ensure_ascii=False)
            assert all(phrase in report_text for phrase in args.expect), "Missing expected content"
            assert output["findings"], "No grounded findings"
            assert all(finding["source_id"] == source["id"] for finding in output["findings"])
            await page.goto(url + "#/give/" + task["id"], wait_until="domcontentloaded")
            button = page.locator("#sourceDownloads [data-artifact]").filter(has_text="DOCX")
            await button.wait_for(timeout=30000)
            async with page.expect_download() as pending_download:
                await button.click()
            download = await pending_download.value
            report = args.output / "report.docx"
            await download.save_as(str(report))
            assert report.stat().st_size > 1000
            document = Document(str(report))
            downloaded_text = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [
                    cell.text
                    for table in document.tables
                    for row in table.rows
                    for cell in row.cells
                ]
            )
            assert output["summary"] in downloaded_text
            assert all(finding["quote"] in downloaded_text for finding in output["findings"])
            receipt["download_bytes"] = report.stat().st_size
            receipt["download_content_matches"] = True
            receipt["findings"] = len(output["findings"])
            receipt["passed"] = not receipt["page_errors"]
        except Exception as exc:
            receipt["exception"] = str(exc)[:2000]
        finally:
            await page.screenshot(path=str(args.output / "result.png"), full_page=True)
            (args.output / "receipt.json").write_text(
                json.dumps(receipt, indent=2, ensure_ascii=False)
            )
            await browser.close()
    print(json.dumps(receipt, ensure_ascii=False))
    if not receipt["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base", default="http://127.0.0.1:8100")
    parser.add_argument("--org", required=True)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--goal", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--chromium", required=True)
    parser.add_argument("--expect", action="append", default=[])
    parser.add_argument("--timeout", type=int, default=300)
    asyncio.run(main(parser.parse_args()))
