"""Source validation, archive boundaries and exported Office files, read back for real."""

import copy
import io
import json
import zipfile

import pytest
from docx import Document
from openpyxl import load_workbook

from ai_orchestrator.application.source_exports import render_exports
from ai_orchestrator.domain.source_report import compare_quotes, validate_report
from ai_orchestrator.ingest.source_files import checked_zip, extract

CSV = (
    b"quote_id,item,quantity,unit_price,delivery,certified\n"
    b"A,Valve,12,120000,80000,true\n"
    b"B,Valve,12,110000,0,false\n"
    b"C,Valve,12,130000,0,true\n"
)


def source():
    parsed = extract("quotes.csv", CSV)
    return {"id": "source-A", **parsed}


def report():
    item = source()
    return {
        "summary": "A is the lowest eligible quotation; no order has been placed.",
        "findings": [
            {
                "text": "Certification is required; B is ineligible.",
                "source_id": item["id"],
                "locator": item["segments"][0]["locator"],
                "quote": "A,Valve,12,120000,80000,true",
            }
        ],
        "limitations": ["Synthetic source; no external supplier verification."],
        "winner": "A",
        "total_vnd": 1520000,
    }


def test_exact_money_eligibility_and_grounding():
    item = source()
    expected = compare_quotes(item["tables"])
    assert expected == {"winner": "A", "total_vnd": 1520000}
    validate_report(report(), [item], expected)


@pytest.mark.parametrize(
    "field,value", [("winner", "B"), ("total_vnd", 1320000), ("total_vnd", True)]
)
def test_plausible_but_wrong_business_output_is_rejected(field, value):
    output = report()
    output[field] = value
    with pytest.raises(ValueError):
        validate_report(output, [source()], {"winner": "A", "total_vnd": 1520000})


@pytest.mark.parametrize(
    "field,value",
    [
        ("source_id", "another-tenant-file"),
        ("locator", "made-up-page"),
        ("quote", "Fabricated certificate is valid"),
    ],
)
def test_invented_evidence_cannot_complete_a_report(field, value):
    output = report()
    output["findings"][0][field] = value
    with pytest.raises(ValueError):
        validate_report(output, [source()], {"winner": "A", "total_vnd": 1520000})


@pytest.mark.parametrize(
    "key,value",
    [
        ("unit_price", "NaN"),
        ("quantity", "0"),
        ("quantity", "1.5"),
        ("delivery", "-1"),
        ("certified", ""),
        ("item", "Different-unit"),
    ],
)
def test_malformed_or_incomparable_quotes_are_not_calculated(key, value):
    tables = copy.deepcopy(source()["tables"])
    tables[0]["rows"][0][key] = value
    with pytest.raises(ValueError):
        compare_quotes(tables)


def test_none_eligible_is_an_explicit_no_winner_result():
    tables = copy.deepcopy(source()["tables"])
    for row in tables[0]["rows"]:
        row["certified"] = "false"
    assert compare_quotes(tables) == {"winner": None, "total_vnd": None}


@pytest.mark.parametrize(
    "filename", ["../../outside.txt", "/tmp/outside.txt", "a\\..\\outside.txt", "C:/outside.txt"]
)
def test_zip_paths_cannot_escape_storage(filename):
    packet = io.BytesIO()
    with zipfile.ZipFile(packet, "w") as archive:
        archive.writestr(filename, "unsafe content")
    with pytest.raises(ValueError):
        checked_zip(packet.getvalue())


def test_zip_bomb_is_rejected_before_member_read():
    packet = io.BytesIO()
    with zipfile.ZipFile(packet, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("huge.txt", b"0" * 1000000)
    with pytest.raises(ValueError, match="compression"):
        checked_zip(packet.getvalue())


def test_nested_zip_retains_member_provenance():
    inner = io.BytesIO()
    with zipfile.ZipFile(inner, "w") as archive:
        archive.writestr("quotes.csv", CSV)
    outer = io.BytesIO()
    with zipfile.ZipFile(outer, "w") as archive:
        archive.writestr("bundle.zip", inner.getvalue())
    result = extract("sources.zip", outer.getvalue())
    assert result["state"] == "ready"
    assert result["segments"][0]["locator"].startswith("bundle.zip#quotes.csv#")
    assert compare_quotes(result["tables"])["winner"] == "A"


def test_unknown_format_and_empty_sources_are_explicit():
    assert extract("drawing.dwg", b"opaque binary")["state"] == "unsupported"
    assert extract("empty.txt", b" \n")["state"] == "needs_review"
    with pytest.raises(ValueError):
        extract("fake.csv", b"a\x00binary")


def test_real_office_exports_open_and_escape_formulas():
    output = report()
    output["findings"][0]["text"] = '=HYPERLINK("https://example.invalid", "do not execute")'
    manifest = {
        "sha256": "a" * 64,
        "sources": [
            {
                "id": "source-A",
                "filename": "quotes.csv",
                "sha256": "b" * 64,
                "classification": "internal",
                "synthetic": True,
            }
        ],
    }
    exports = render_exports("So sánh báo giá tiếng Việt", output, manifest)
    doc = Document(io.BytesIO(exports["docx"][0]))
    assert "MOCK" in "\n".join(p.text for p in doc.paragraphs)
    assert output["findings"][0]["quote"] in doc.tables[0].rows[1].cells[2].text
    book = load_workbook(io.BytesIO(exports["xlsx"][0]), data_only=False)
    cells = [
        cell for row in book.active for cell in row if cell.value == output["findings"][0]["text"]
    ]
    assert cells and cells[0].data_type == "s"
    with zipfile.ZipFile(io.BytesIO(exports["zip"][0])) as archive:
        assert {"report.docx", "report.xlsx", "report.json", "report.md", "manifest.json"} <= set(
            archive.namelist()
        )
        assert json.loads(archive.read("report.json"))["output"]["total_vnd"] == 1520000
    parsed = extract("report.docx", exports["docx"][0])
    assert parsed["state"] == "ready"
    assert any("tiếng Việt" in s["text"] for s in parsed["segments"])


def test_report_must_not_mix_in_a_foreign_script():
    """The live failure: an exact Vietnamese quote, a Chinese word in the text.

    A quote can be a verbatim substring while the prose around it leaks a
    different script, so the surface guard reads the finding text, not just the
    citation.
    """
    source_item = {
        "id": "src_hr",
        "segments": [
            {
                "locator": "lines:1-2",
                "text": "HR Agent tối đa L2 (sàng lọc, soạn thảo, nhắc quy trình). "
                "Vùng cấm AI: mọi quyết định tuyển/thôi việc/kỷ luật/lương do con người.",
            }
        ],
    }
    output = {
        "summary": "Bản tổng hợp HR giữ nguyên giới hạn quyền của trợ lý.",
        "findings": [
            {
                "text": "HR Agent tối đa L2 với chức năng筛查, soạn thảo, nhắc quy trình.",
                "source_id": "src_hr",
                "locator": "lines:1-2",
                "quote": "HR Agent tối đa L2 (sàng lọc, soạn thảo, nhắc quy trình)",
            }
        ],
        "limitations": ["Chỉ dùng nguồn được chọn."],
    }
    with pytest.raises(ValueError, match="script absent from the sources"):
        validate_report(output, [source_item])
    # The same finding written in the source's own Vietnamese passes.
    output["findings"][0]["text"] = (
        "HR Agent tối đa L2 với chức năng sàng lọc, soạn thảo, nhắc quy trình."
    )
    validate_report(output, [source_item])


def test_foreign_script_is_allowed_when_the_source_contains_it():
    from ai_orchestrator.domain.source_report import foreign_script_leakage

    assert foreign_script_leakage("筛查", set()) == ["查", "筛"]
    # A genuinely multilingual source is not misflagged.
    assert foreign_script_leakage("筛查", {"筛", "查"}) == []
    assert foreign_script_leakage("sàng lọc — 30-60-90", set()) == []
