"""Render accepted report data to real files; no model-generated binary content."""

from __future__ import annotations

import io
import json
import zipfile
from html import escape
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import Workbook
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

pdfmetrics.registerFont(
    TTFont(
        "SourceReport",
        str(Path(__file__).resolve().parents[1] / "report_assets" / "DejaVuSans.ttf"),
    )
)


def render_exports(
    title: str,
    output: dict[str, Any],
    manifest: dict[str, Any],
    originals: list[tuple[str, bytes]] | None = None,
) -> dict[str, tuple[bytes, str]]:
    label = (
        "CONTAINS MOCK SOURCES / CÓ NGUỒN MOCK"
        if any(s["synthetic"] for s in manifest["sources"])
        else "DRAFT FOR REVIEW / BẢN NHÁP ĐỂ RÀ SOÁT"
    )
    lines = [title, label, output["summary"]]
    if "winner" in output:
        lines.append(f"Winner: {output['winner']} | Total VND: {output['total_vnd']}")
    doc = Document()
    doc.add_heading(title, 0)
    doc.add_paragraph(label)
    doc.add_paragraph(output["summary"])
    if "winner" in output:
        doc.add_paragraph(lines[-1])
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(
        table.rows[0].cells, ("Finding", "Source / location", "Verbatim evidence"), strict=True
    ):
        cell.text = text
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Report"

    def row(values: list[Any]) -> None:
        sheet.append(values)
        # Uploaded/model text must remain literal text, never spreadsheet formulas.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    row([title, label])
    row(["Summary", output["summary"]])
    if "winner" in output:
        row(["Winner", output["winner"], "Total VND", output["total_vnd"]])
    row(["Finding", "Source", "Locator", "Quote"])
    for finding in output["findings"]:
        cells = table.add_row().cells
        cells[0].text = finding["text"]
        cells[1].text = finding["source_id"] + " / " + finding["locator"]
        cells[2].text = finding["quote"]
        row([finding["text"], finding["source_id"], finding["locator"], finding["quote"]])
        lines.append(
            f"- {finding['text']}\n  "
            f"[{finding['source_id']} / {finding['locator']}] {finding['quote']}"
        )
    doc.add_heading("Limitations / Giới hạn", level=1)
    for limitation in output["limitations"]:
        doc.add_paragraph(limitation)
        row(["Limitation", limitation])
        lines.append("- " + limitation)
    doc.add_heading("Source manifest", level=1)
    for source in manifest["sources"]:
        text = (
            f"{source['id']} | {source['filename']} | SHA-256 {source['sha256']} | "
            f"{source['classification']} | mock={source['synthetic']}"
        )
        doc.add_paragraph(text)
        row(["Source manifest", text])
        lines.append(text)
    doc.add_paragraph("Manifest SHA-256: " + manifest["sha256"])
    row(["Manifest SHA-256", manifest["sha256"]])
    sheet.freeze_panes = "A5"
    for column, width in [("A", 55), ("B", 45), ("C", 28), ("D", 80)]:
        sheet.column_dimensions[column].width = width
    docx, xlsx = io.BytesIO(), io.BytesIO()
    doc.save(docx)
    workbook.save(xlsx)
    pdf = io.BytesIO()
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = "SourceReport"
        style.alignment = TA_LEFT
        style.wordWrap = "CJK"
    story = []
    for index, value in enumerate(lines):
        story.extend(
            [
                Paragraph(
                    escape(value).replace("\n", "<br/>"),
                    styles["Heading1"] if index == 0 else styles["BodyText"],
                ),
                Spacer(1, 8),
            ]
        )
    SimpleDocTemplate(pdf, pagesize=A4, title=title).build(story)
    exports = {
        "pdf": (pdf.getvalue(), "application/pdf"),
        "docx": (
            docx.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        "xlsx": (
            xlsx.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
        "json": (
            json.dumps(
                {"output": output, "manifest": manifest}, ensure_ascii=False, indent=2
            ).encode(),
            "application/json",
        ),
        "md": ("\n\n".join(lines).encode(), "text/markdown; charset=utf-8"),
    }
    packet = io.BytesIO()
    with zipfile.ZipFile(packet, "w", zipfile.ZIP_DEFLATED) as archive:
        for extension, (content, _) in exports.items():
            archive.writestr("report." + extension, content)
        archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
        for name, content in originals or []:
            archive.writestr("sources/" + name, content)
    exports["zip"] = (packet.getvalue(), "application/zip")
    return exports
