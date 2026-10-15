"""Bounded, offline extraction. Invoked only inside the resource-limited worker."""

from __future__ import annotations

import csv
import io
import json
import re
import shutil
import stat
import subprocess
import tempfile
import zipfile
from pathlib import Path, PurePosixPath
from typing import Any

from defusedxml import ElementTree as ET

MAX_BYTES = 10 * 1024 * 1024
MAX_EXPANDED = 24 * 1024 * 1024
MAX_TEXT = 60_000
MEDIA = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".csv": "text/csv",
    ".json": "application/json",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".zip": "application/zip",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
}


def _native(command: list[str], **kwargs: Any) -> subprocess.CompletedProcess[bytes]:
    binary = shutil.which(command[0])
    if binary is None:
        raise ValueError("Required extraction tool unavailable: " + command[0])
    # argv contains only a fixed tool and locally generated temporary paths; no shell.
    return subprocess.run([binary, *command[1:]], **kwargs)  # noqa: S603


def checked_zip(data: bytes) -> zipfile.ZipFile:
    archive = zipfile.ZipFile(io.BytesIO(data))
    entries = archive.infolist()
    if len(entries) > 300 or sum(x.file_size for x in entries) > MAX_EXPANDED:
        raise ValueError("Archive exceeds entry or expanded-size limit")
    for entry in entries:
        path = PurePosixPath(entry.filename)
        if (
            path.is_absolute()
            or ".." in path.parts
            or "\\" in entry.filename
            or ":" in entry.filename
            or "\x00" in entry.filename
            or stat.S_ISLNK(entry.external_attr >> 16)
            or entry.flag_bits & 1
        ):
            raise ValueError("Archive contains unsafe paths, links or encryption")
        if entry.file_size > MAX_BYTES or entry.file_size > max(1, entry.compress_size) * 200:
            raise ValueError("Archive member exceeds size or compression-ratio limit")
    return archive


def extract(filename: str, data: bytes, depth: int = 0) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    result: dict[str, Any] = {
        "parser": "source-v1",
        "state": "ready",
        "segments": [],
        "tables": [],
        "warnings": [],
    }
    segments = result["segments"]

    def add(locator: str, value: str) -> None:
        value = value.strip()
        if value:
            segments.append({"locator": locator, "text": value})

    if not data or len(data) > MAX_BYTES:
        raise ValueError("Empty file or file exceeds 10 MiB")
    if suffix in {".txt", ".md", ".csv", ".json"}:
        text = data.decode("utf-8-sig")
        if "\x00" in text:
            raise ValueError("Text file contains binary NUL bytes")
        if suffix == ".json":
            json.loads(text)
        for start in range(0, len(text.splitlines()), 12):
            add(
                f"lines:{start + 1}-{min(start + 12, len(text.splitlines()))}",
                "\n".join(text.splitlines()[start : start + 12]),
            )
        if suffix == ".csv":
            rows = list(csv.DictReader(io.StringIO(text)))
            if len(rows) > 5000:
                raise ValueError("CSV exceeds row limit")
            result["tables"] = [{"locator": "rows", "rows": rows}]
    elif suffix in {".docx", ".pptx", ".xlsx", ".zip"}:
        with checked_zip(data) as archive:
            if suffix != ".zip" and any(
                name.lower().endswith("vbaproject.bin") for name in archive.namelist()
            ):
                raise ValueError("Macro-bearing Office documents are not supported")
            if suffix == ".docx":
                document = ET.fromstring(archive.read("word/document.xml"))
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                for index, paragraph in enumerate(document.findall(".//w:p", ns), 1):
                    add(
                        f"paragraph:{index}",
                        "".join(n.text or "" for n in paragraph.findall(".//w:t", ns)),
                    )
            elif suffix == ".pptx":
                for name in sorted(archive.namelist()):
                    if re.fullmatch(r"ppt/slides/slide\d+\.xml", name):
                        root = ET.fromstring(archive.read(name))
                        add(
                            name,
                            " ".join(n.text or "" for n in root.iter() if n.tag.endswith("}t")),
                        )
            elif suffix == ".xlsx":
                from openpyxl import load_workbook

                book = load_workbook(
                    io.BytesIO(data), read_only=True, data_only=False, keep_links=False
                )
                try:
                    count = 0
                    has_formulas = False
                    for sheet in book:
                        for row in sheet.iter_rows():
                            count += 1
                            if count > 5000:
                                raise ValueError("Workbook exceeds row limit")
                            has_formulas |= any(cell.data_type == "f" for cell in row)
                            cells = [
                                f"{cell.coordinate}={cell.value}"
                                for cell in row
                                if cell.value is not None
                            ]
                            add(f"{sheet.title}!row:{row[0].row}", " | ".join(cells))
                    if has_formulas:
                        result["state"] = "needs_review"
                    result["warnings"].append(
                        "Formulas are preserved as formulas; "
                        "missing cached values are not treated as zero."
                    )
                finally:
                    book.close()
            else:
                if depth >= 2 or len(archive.infolist()) > 30:
                    raise ValueError("ZIP exceeds nesting or member limit")
                for entry in archive.infolist():
                    if entry.is_dir():
                        continue
                    child = extract(entry.filename, archive.read(entry), depth + 1)
                    if child["state"] != "ready":
                        result["state"] = "needs_review"
                        result["warnings"].append(f"{entry.filename}: {child['state']}")
                    for segment in child["segments"]:
                        add(entry.filename + "#" + segment["locator"], segment["text"])
                    for table in child["tables"]:
                        result["tables"].append(
                            {**table, "locator": entry.filename + "#" + table["locator"]}
                        )
                    result["warnings"].extend(child["warnings"])
    elif suffix == ".pdf" or suffix in {".jpg", ".jpeg", ".png", ".tif", ".tiff"}:
        with tempfile.TemporaryDirectory(prefix="ao-extract-") as folder:
            path = Path(folder) / ("input" + suffix)
            path.write_bytes(data)
            if suffix == ".pdf":
                if not data.startswith(b"%PDF-"):
                    raise ValueError("File is not a PDF")
                if not shutil.which("pdftotext"):
                    raise ValueError("PDF extraction requires pdftotext on the worker")
                raw = _native(
                    ["pdftotext", "-layout", str(path), "-"],
                    check=True,
                    capture_output=True,
                    timeout=15,
                ).stdout.decode("utf-8")
                pages = raw.split("\f")
                if len(pages) > 51:
                    raise ValueError("PDF exceeds 50-page limit")
                if len(raw.strip()) >= 20:
                    for index, page in enumerate(pages, 1):
                        add(f"page:{index}", page)
                    # Mixed scan/text PDFs must not silently drop scanned pages.
                    if any(not p.strip() for p in pages[:-1]):
                        result["state"] = "needs_review"
                        result["warnings"].append(
                            "PDF contains an empty or scanned page; review extraction before use."
                        )
                else:
                    if not shutil.which("pdftoppm"):
                        raise ValueError("Scanned PDF requires pdftoppm and OCR")
                    _native(
                        [
                            "pdftoppm",
                            "-f",
                            "1",
                            "-l",
                            "10",
                            "-scale-to",
                            "1800",
                            "-png",
                            str(path),
                            str(Path(folder) / "page"),
                        ],
                        check=True,
                        capture_output=True,
                        timeout=15,
                    )
                    for image in sorted(Path(folder).glob("page-*.png")):
                        text, confidence = ocr(image)
                        add("page:" + str(int(image.stem.split("-")[-1])), text)
                        result["warnings"].append(
                            f"OCR confidence {confidence:.1f}; review scanned source."
                        )
                    result["state"] = "needs_review"
            else:
                from PIL import Image

                with Image.open(path) as decoded_image:
                    expected = {
                        ".jpg": "JPEG",
                        ".jpeg": "JPEG",
                        ".png": "PNG",
                        ".tif": "TIFF",
                        ".tiff": "TIFF",
                    }
                    if (
                        decoded_image.format != expected[suffix]
                        or decoded_image.width * decoded_image.height > 20_000_000
                    ):
                        raise ValueError("Image format mismatch or image exceeds 20 million pixels")
                    if getattr(decoded_image, "n_frames", 1) > 1:
                        raise ValueError("Multi-frame images require splitting before upload")
                    decoded_image.verify()
                text, confidence = ocr(path)
                add("image:1", text)
                result["warnings"].append(f"OCR mean confidence {confidence:.1f}")
                if confidence < 70:
                    result["state"] = "needs_review"
    else:
        result["state"] = "unsupported"
        result["warnings"].append("This format can be stored but is not extractable yet.")
    length = sum(len(s["text"]) for s in segments)
    if length > MAX_TEXT:
        raise ValueError("Extracted text exceeds 60000 characters; split the source")
    if not segments and result["state"] == "ready":
        result["state"] = "needs_review"
        result["warnings"].append("No readable content found; supply another source.")
    return result


def ocr(path: Path) -> tuple[str, float]:
    if not shutil.which("tesseract"):
        raise ValueError("OCR requires tesseract with eng and vie language data")
    raw = _native(
        ["tesseract", str(path), "stdout", "-l", "eng+vie", "tsv"],
        check=True,
        capture_output=True,
        timeout=20,
    ).stdout.decode("utf-8")
    rows = list(csv.DictReader(io.StringIO(raw), delimiter="\t"))
    words = [r for r in rows if (r.get("text") or "").strip()]
    confidence = sum(float(r["conf"]) for r in words) / len(words) if words else 0.0
    return " ".join(r["text"] for r in words), confidence
